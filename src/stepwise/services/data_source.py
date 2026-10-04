"""Data file loader for Excel (.xlsx) and CSV files (Section 10.1).

Features:
- openpyxl with read_only=True and data_only=True
- Auto-detect CSV encoding (utf-8-sig, cp949, cp1250, etc.) via charset-normalizer
- Auto-detect CSV delimiter (comma, semicolon, tab)
- Strict value normalization (strip unnecessary '.0' on integers, ISO date format)
- Skips completely blank rows and returns original 1-indexed row numbers
"""

from __future__ import annotations

import csv
import datetime
import os
from typing import Any

import charset_normalizer
import openpyxl


def normalize_cell_value(val: Any) -> str:
    """Normalize cell values to clean strings following Section 10.1."""
    if val is None:
        return ""

    if isinstance(val, (datetime.date, datetime.datetime)):
        # Default date format YYYY-MM-DD
        return val.strftime("%Y-%m-%d")

    if isinstance(val, float):
        # Remove redundant '.0' for integer values
        if val.is_integer():
            return str(int(val))
        return str(val)

    if isinstance(val, int):
        return str(val)

    text = str(val).strip()
    # Check if text looks like "1200.0"
    if text.endswith(".0"):
        try:
            f = float(text)
            if f.is_integer():
                return str(int(f))
        except ValueError:
            pass

    return text


def detect_csv_encoding_and_delimiter(
    filepath: str,
    override_encoding: str | None = None,
    override_delimiter: str | None = None,
) -> tuple[str, str]:
    """Detect text encoding and delimiter for a CSV file."""
    encoding = override_encoding
    if not encoding:
        with open(filepath, "rb") as f:
            raw_sample = f.read(32768)
        # Check BOM first
        if raw_sample.startswith(b"\xef\xbb\xbf"):
            encoding = "utf-8-sig"
        else:
            matches = charset_normalizer.from_bytes(raw_sample)
            best = matches.best()
            encoding = best.encoding if best else "utf-8"

    # Detect delimiter
    delimiter = override_delimiter
    if not delimiter:
        try:
            with open(filepath, encoding=encoding, errors="replace") as f:
                sample_lines = [f.readline() for _ in range(5)]
            sample_text = "".join(sample_lines)

            sniffer = csv.Sniffer()
            dialect = sniffer.sniff(sample_text, delimiters=";,\t|")
            delimiter = dialect.delimiter
        except Exception:
            delimiter = ","

    return encoding, delimiter


def read_data_file(
    filepath: str,
    sheet: str | None = None,
    header_row: int = 1,
    encoding: str | None = None,
    delimiter: str | None = None,
) -> tuple[list[str], list[dict[str, str]], list[int]]:
    """Read .xlsx or .csv data file.

    Returns:
    - headers: list of column name strings
    - rows: list of row dicts {column_name: string_value}
    - row_numbers: list of 1-indexed source row numbers
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"Data file not found: {filepath}")

    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".xlsx":
        return _read_xlsx(filepath, sheet=sheet, header_row=header_row)
    else:
        return _read_csv(
            filepath,
            header_row=header_row,
            override_encoding=encoding,
            override_delimiter=delimiter,
        )


def _read_xlsx(
    filepath: str,
    sheet: str | None,
    header_row: int,
) -> tuple[list[str], list[dict[str, str]], list[int]]:
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        if ws is None:
            raise ValueError(f"No active sheet in workbook {filepath}")

        headers: list[str] = []
        rows: list[dict[str, str]] = []
        row_numbers: list[int] = []

        curr_row_idx = 0
        for raw_row in ws.iter_rows(values_only=True):
            curr_row_idx += 1
            if curr_row_idx < header_row:
                continue

            row_vals = [normalize_cell_value(cell) for cell in raw_row]

            if curr_row_idx == header_row:
                # Header row
                headers = [h or f"Column_{i + 1}" for i, h in enumerate(row_vals)]
                continue

            # Check if entire row is blank
            if not any(row_vals):
                continue

            row_dict: dict[str, str] = {}
            for col_idx, col_name in enumerate(headers):
                row_dict[col_name] = row_vals[col_idx] if col_idx < len(row_vals) else ""

            rows.append(row_dict)
            row_numbers.append(curr_row_idx)

        return headers, rows, row_numbers
    finally:
        wb.close()


def _read_csv(
    filepath: str,
    header_row: int,
    override_encoding: str | None,
    override_delimiter: str | None,
) -> tuple[list[str], list[dict[str, str]], list[int]]:
    enc, delim = detect_csv_encoding_and_delimiter(
        filepath, override_encoding=override_encoding, override_delimiter=override_delimiter
    )

    headers: list[str] = []
    rows: list[dict[str, str]] = []
    row_numbers: list[int] = []

    with open(filepath, encoding=enc, errors="replace", newline="") as f:
        reader = csv.reader(f, delimiter=delim)
        curr_row_idx = 0
        for raw_row in reader:
            curr_row_idx += 1
            if curr_row_idx < header_row:
                continue

            row_vals = [normalize_cell_value(cell) for cell in raw_row]

            if curr_row_idx == header_row:
                headers = [h or f"Column_{i + 1}" for i, h in enumerate(row_vals)]
                continue

            if not any(row_vals):
                continue

            row_dict: dict[str, str] = {}
            for col_idx, col_name in enumerate(headers):
                row_dict[col_name] = row_vals[col_idx] if col_idx < len(row_vals) else ""

            rows.append(row_dict)
            row_numbers.append(curr_row_idx)

    return headers, rows, row_numbers
